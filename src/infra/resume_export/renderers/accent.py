from dataclasses import dataclass
from html import escape
from io import BytesIO

from docx import Document
from docx.document import Document as WordDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT, TA_RIGHT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfgen.canvas import Canvas
from reportlab.platypus import HRFlowable, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
from reportlab.platypus.doctemplate import BaseDocTemplate
from reportlab.platypus.flowables import Flowable

from core.i18n.enums import LanguageEnum
from core.resumes.enums import ResumeCurrentStatusEnum
from core.resumes.schemas import (
    ResumeAdditionalSection,
    ResumeCertificationItem,
    ResumeContent,
    ResumeEducationItem,
    ResumeExperienceItem,
    ResumeExportParams,
    ResumeProjectItem,
)
from infra.resume_export.renderers.simple import ResumeExportLabels, SimpleResumeRenderer

INK = colors.HexColor("#1E293B")
BLUE = colors.HexColor("#255487")
MUTED = colors.HexColor("#526174")
PALE = colors.HexColor("#F1F7FC")
PALE_GRAY = colors.HexColor("#F6F9FC")
BORDER = colors.HexColor("#C9D9E8")

PDF_TEXT_STYLES = {
    "name": (22, 26, INK, True, 1, False, TA_LEFT),
    "role": (11, 14, BLUE, False, 3, False, TA_LEFT),
    "location": (9, 12, MUTED, False, 3, False, TA_LEFT),
    "contact": (8.5, 11, INK, False, 0, False, TA_LEFT),
    "section": (10.5, 13, BLUE, True, 1, True, TA_LEFT),
    "body": (8.8, 11.3, INK, False, 1.5, False, TA_LEFT),
    "bold": (8.8, 11.3, INK, True, 1.5, False, TA_LEFT),
    "heading": (8.8, 11.3, INK, True, 1.5, True, TA_LEFT),
    "bullet": (8.8, 11.3, INK, False, 0.5, False, TA_LEFT),
    "muted": (8.8, 11.3, MUTED, False, 1.5, False, TA_LEFT),
    "band": (8.2, 10.5, MUTED, False, 0, False, TA_LEFT),
    "company": (10, 11.3, INK, True, 0, False, TA_LEFT),
    "company_meta": (8.5, 10.5, MUTED, False, 0, False, TA_LEFT),
    "company_period": (8.5, 10.5, MUTED, False, 0, False, TA_RIGHT),
    "project": (9, 11.3, INK, True, 0, False, TA_LEFT),
    "link": (8.8, 11.3, BLUE, False, 1.5, False, TA_LEFT),
}


@dataclass(frozen=True, slots=True, kw_only=True)
class AccentResumeRenderer(SimpleResumeRenderer):
    def _export_pdf(self, *, params: ResumeExportParams) -> bytes:
        self._register_pdf_fonts()
        labels = self._labels_for_language(language=params.language)
        output = BytesIO()
        document = SimpleDocTemplate(
            output,
            pagesize=A4,
            rightMargin=15 * mm,
            leftMargin=15 * mm,
            topMargin=16 * mm,
            bottomMargin=18 * mm,
            title=params.title,
            author=params.content.profile.full_name,
        )
        width = A4[0] - 30 * mm
        story: list[Flowable] = []
        self._pdf_profile(
            story=story, content=params.content, language=params.language, width=width
        )
        self._pdf_introduction(story=story, content=params.content, labels=labels, width=width)
        self._pdf_experience_section(
            story=story, content=params.content, labels=labels, width=width
        )
        self._pdf_other_sections(story=story, content=params.content, labels=labels, width=width)

        def footer(canvas: Canvas, doc: BaseDocTemplate) -> None:
            canvas.saveState()
            canvas.setStrokeColor(BORDER)
            canvas.line(15 * mm, 14 * mm, A4[0] - 15 * mm, 14 * mm)
            canvas.setFillColor(MUTED)
            canvas.setFont(self.font_regular_name, 8)
            canvas.drawString(
                15 * mm, 10 * mm, "Резюме" if params.language == LanguageEnum.RU else "Resume"
            )
            page_label = "Страница" if params.language == LanguageEnum.RU else "Page"
            canvas.drawRightString(A4[0] - 15 * mm, 10 * mm, f"{page_label} {doc.page}")
            canvas.restoreState()

        document.build(story, onFirstPage=footer, onLaterPages=footer)
        return output.getvalue()

    def _pdf_introduction(
        self,
        *,
        story: list[Flowable],
        content: ResumeContent,
        labels: ResumeExportLabels,
        width: float,
    ) -> None:
        if content.summary.text:
            self._pdf_section(story=story, title=labels.summary, width=width)
            story.append(self._pdf_paragraph(content.summary.text, kind="body"))
        if content.skills:
            self._pdf_section(story=story, title=labels.skills, width=width)
            story.extend(
                self._pdf_paragraph(f"{group.category}: {', '.join(group.items)}", kind="body")
                for group in content.skills
            )
        if content.education:
            self._pdf_section(story=story, title=labels.education, width=width)
            for education in content.education:
                self._pdf_education(story=story, item=education, labels=labels)
        if content.languages:
            self._pdf_section(story=story, title=labels.languages, width=width)
            story.extend(
                self._pdf_paragraph(f"{language.name} | {language.proficiency}", kind="body")
                for language in content.languages
            )

    def _pdf_experience_section(
        self,
        *,
        story: list[Flowable],
        content: ResumeContent,
        labels: ResumeExportLabels,
        width: float,
    ) -> None:
        if not content.experience:
            return
        self._pdf_section(story=story, title=labels.experience, width=width)
        for experience in content.experience:
            self._pdf_experience(story=story, item=experience, labels=labels, width=width)

    def _pdf_other_sections(
        self,
        *,
        story: list[Flowable],
        content: ResumeContent,
        labels: ResumeExportLabels,
        width: float,
    ) -> None:
        if content.certifications:
            self._pdf_section(story=story, title=labels.certifications, width=width)
            for certification in content.certifications:
                story.append(
                    self._pdf_paragraph(
                        f"{certification.name} | {certification.issuer}", kind="bold"
                    )
                )
                dates = self._certification_dates(item=certification, language=labels.language)
                if dates:
                    story.append(self._pdf_paragraph(dates, kind="muted"))
                if certification.credential_url:
                    story.append(self._pdf_paragraph(certification.credential_url, kind="link"))
        for extra in content.additional_sections:
            self._pdf_additional(story=story, section=extra, width=width)

    def _pdf_profile(
        self, *, story: list[Flowable], content: ResumeContent, language: LanguageEnum, width: float
    ) -> None:
        profile = content.profile
        story.append(self._pdf_paragraph(profile.full_name, kind="name"))
        story.append(self._pdf_paragraph(profile.role, kind="role"))
        if profile.location:
            story.append(self._pdf_paragraph(profile.location, kind="location"))
        contacts = [
            ("Телефон" if language == LanguageEnum.RU else "Phone", profile.phone),
            ("Email", profile.email),
            ("Telegram", profile.telegram),
            ("LinkedIn", profile.linkedin_url),
            ("GitHub", profile.github_url),
            ("Сайт" if language == LanguageEnum.RU else "Website", profile.website_url),
        ]
        contacts = [(label, value) for label, value in contacts if value]
        if contacts:
            heading = "КОНТАКТНЫЕ ДАННЫЕ" if language == LanguageEnum.RU else "CONTACTS"
            self._pdf_section(story=story, title=heading, width=width)
            rows = []
            for index in range(0, len(contacts), 2):
                pair = contacts[index : index + 2]
                rows.append(
                    [
                        self._pdf_paragraph(f"{label}: {value}", kind="contact")
                        for label, value in pair
                    ]
                    + [""] * (2 - len(pair))
                )
            table = Table(rows, colWidths=[width / 2, width / 2], hAlign="LEFT")
            table.setStyle(
                TableStyle(
                    [
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("LEFTPADDING", (0, 0), (-1, -1), 0),
                        ("RIGHTPADDING", (0, 0), (-1, -1), 8),
                        ("TOPPADDING", (0, 0), (-1, -1), 1),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 1),
                    ]
                )
            )
            story.append(table)

    def _pdf_section(self, *, story: list[Flowable], title: str, width: float) -> None:
        story.append(Spacer(1, 6))
        story.append(self._pdf_paragraph(title.upper(), kind="section"))
        story.append(HRFlowable(width=width, thickness=0.7, color=BLUE, spaceAfter=3))

    def _pdf_paragraph(self, text: str, *, kind: str) -> Paragraph:
        size, leading, color, bold, after, keep, align = PDF_TEXT_STYLES[kind]
        style = ParagraphStyle(
            name="AccentText",
            fontName=self.font_bold_name if bold else self.font_regular_name,
            fontSize=size,
            leading=leading,
            textColor=color,
            spaceAfter=after,
            alignment=align,
            keepWithNext=keep,
        )
        return Paragraph(escape(text, quote=False).replace("\n", "<br/>"), style)

    def _pdf_bullets(self, *, story: list[Flowable], items: list[str]) -> None:
        story.extend(self._pdf_paragraph(f"•  {item}", kind="bullet") for item in items)

    def _pdf_band(self, *, story: list[Flowable], text: str, width: float) -> None:
        if not text:
            return
        table = Table(
            [[self._pdf_paragraph(text, kind="band")]],
            colWidths=[width],
            hAlign="LEFT",
        )
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, -1), PALE_GRAY),
                    ("LEFTPADDING", (0, 0), (-1, -1), 5),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                    ("TOPPADDING", (0, 0), (-1, -1), 3),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                ]
            )
        )
        story.append(table)

    def _pdf_experience(
        self,
        *,
        story: list[Flowable],
        item: ResumeExperienceItem,
        labels: ResumeExportLabels,
        width: float,
    ) -> None:
        story.append(Spacer(1, 5))
        period = self._format_date_range(
            start_date=item.start_date,
            end_date=item.end_date,
            current_status=item.current_status,
            labels=labels,
        )
        card_rows = [
            [
                self._pdf_paragraph(item.company, kind="company"),
                self._pdf_paragraph(period, kind="company_period"),
            ],
            [
                self._pdf_paragraph(
                    " | ".join(part for part in (item.position, item.location) if part),
                    kind="company_meta",
                ),
                "",
            ],
        ]
        if item.summary:
            card_rows.append([self._pdf_paragraph(item.summary, kind="band"), ""])
        card = Table(card_rows, colWidths=[width * 0.72, width * 0.28], hAlign="LEFT")
        card.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, -1), PALE),
                    ("BOX", (0, 0), (-1, -1), 0.5, BORDER),
                    ("SPAN", (0, 1), (1, 1)),
                    *((("SPAN", (0, 2), (1, 2)),) if item.summary else ()),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 7),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 7),
                    ("TOPPADDING", (0, 0), (-1, -1), 2),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
                ]
            )
        )
        story.append(card)
        self._pdf_bullets(story=story, items=item.highlights)
        if item.technologies:
            self._pdf_band(
                story=story,
                text=f"{labels.technologies}: {', '.join(item.technologies)}",
                width=width,
            )
        for project in item.projects:
            self._pdf_project(
                story=story, project=project, labels=labels, width=width, position=item.position
            )

    def _pdf_project(
        self,
        *,
        story: list[Flowable],
        project: ResumeProjectItem,
        labels: ResumeExportLabels,
        width: float,
        position: str,
    ) -> None:
        story.append(Spacer(1, 4))
        title = f"{labels.project}: {project.name}"
        if project.role and project.role != position:
            title += f" | {project.role}"
        header = Table(
            [[self._pdf_paragraph(title, kind="project")]],
            colWidths=[width],
            hAlign="LEFT",
        )
        header.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, -1), PALE_GRAY),
                    ("LINEBEFORE", (0, 0), (0, -1), 2, BLUE),
                    ("LEFTPADDING", (0, 0), (-1, -1), 7),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                    ("TOPPADDING", (0, 0), (-1, -1), 4),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                ]
            )
        )
        story.append(header)
        self._pdf_bullets(story=story, items=project.highlights)
        self._pdf_band(story=story, text=project.description, width=width)
        if project.technologies:
            self._pdf_band(
                story=story,
                text=f"{labels.technologies}: {', '.join(project.technologies)}",
                width=width,
            )
        if project.url:
            story.append(self._pdf_paragraph(project.url, kind="link"))

    def _pdf_education(
        self, *, story: list[Flowable], item: ResumeEducationItem, labels: ResumeExportLabels
    ) -> None:
        story.append(self._pdf_paragraph(item.institution, kind="bold"))
        story.append(
            self._pdf_paragraph(
                " | ".join(part for part in (item.degree, item.field) if part), kind="body"
            )
        )
        period = self._format_date_range(
            start_date=item.start_date,
            end_date=item.end_date,
            current_status=ResumeCurrentStatusEnum.NOT_SET,
            labels=labels,
        )
        if item.location or period:
            story.append(
                self._pdf_paragraph(
                    " | ".join(part for part in (item.location, period) if part), kind="muted"
                )
            )
        if item.description:
            story.append(self._pdf_paragraph(item.description, kind="body"))

    def _pdf_additional(
        self, *, story: list[Flowable], section: ResumeAdditionalSection, width: float
    ) -> None:
        self._pdf_section(story=story, title=section.title, width=width)
        for item in section.items:
            story.append(self._pdf_paragraph(item.title, kind="heading"))
            if item.description:
                story.append(self._pdf_paragraph(item.description, kind="body"))
            if item.url:
                story.append(self._pdf_paragraph(item.url, kind="link"))

    def _export_docx(self, *, params: ResumeExportParams) -> bytes:
        document = Document()
        section = document.sections[0]
        section.top_margin = Inches(0.65)
        section.bottom_margin = Inches(0.7)
        section.left_margin = Inches(0.65)
        section.right_margin = Inches(0.65)
        normal = document.styles["Normal"]
        normal.font.name = "Arial"
        normal.font.size = Pt(8.7)
        normal.font.color.rgb = RGBColor(30, 41, 59)
        normal.paragraph_format.space_after = Pt(1.5)
        normal.paragraph_format.line_spacing = 1.0
        bullet = document.styles["List Bullet"]
        bullet.paragraph_format.space_after = Pt(0.5)
        bullet.paragraph_format.line_spacing = 1.0
        labels = self._labels_for_language(language=params.language)
        self._word_profile(document=document, content=params.content, language=params.language)
        self._word_introduction(document=document, content=params.content, labels=labels)
        self._word_experience_section(document=document, content=params.content, labels=labels)
        self._word_other_sections(document=document, content=params.content, labels=labels)
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        footer.add_run("Резюме | " if params.language == LanguageEnum.RU else "Resume | ")
        field = OxmlElement("w:fldSimple")
        field.set(qn("w:instr"), "PAGE")
        footer._p.append(field)  # noqa: SLF001
        output = BytesIO()
        document.save(output)
        return output.getvalue()

    def _word_profile(
        self, *, document: WordDocument, content: ResumeContent, language: LanguageEnum
    ) -> None:
        profile = content.profile
        name = document.add_paragraph()
        name.paragraph_format.space_after = Pt(0)
        run = name.add_run(profile.full_name)
        run.bold = True
        run.font.size = Pt(22)
        role = document.add_paragraph()
        role.paragraph_format.space_after = Pt(5)
        run = role.add_run(profile.role)
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(37, 84, 135)
        contact_parts = [
            profile.location,
            profile.phone,
            profile.email,
            profile.telegram,
            profile.linkedin_url,
            profile.github_url,
            profile.website_url,
        ]
        if any(contact_parts):
            self._word_section(
                document, "КОНТАКТНЫЕ ДАННЫЕ" if language == LanguageEnum.RU else "CONTACTS"
            )
            document.add_paragraph(" | ".join(part for part in contact_parts if part))

    def _word_introduction(
        self, *, document: WordDocument, content: ResumeContent, labels: ResumeExportLabels
    ) -> None:
        if content.summary.text:
            self._word_section(document, labels.summary)
            document.add_paragraph(content.summary.text)
        if content.skills:
            self._word_section(document, labels.skills)
            for group in content.skills:
                document.add_paragraph(f"{group.category}: {', '.join(group.items)}")
        if content.education:
            self._word_section(document, labels.education)
            for education in content.education:
                self._word_heading(document, education.institution)
                period = self._format_date_range(
                    start_date=education.start_date,
                    end_date=education.end_date,
                    current_status=ResumeCurrentStatusEnum.NOT_SET,
                    labels=labels,
                )
                document.add_paragraph(
                    " | ".join(
                        part
                        for part in (education.degree, education.field, education.location, period)
                        if part
                    )
                )
                if education.description:
                    document.add_paragraph(education.description)
        if content.languages:
            self._word_section(document, labels.languages)
            for language in content.languages:
                document.add_paragraph(f"{language.name} | {language.proficiency}")

    def _word_experience_section(
        self, *, document: WordDocument, content: ResumeContent, labels: ResumeExportLabels
    ) -> None:
        if not content.experience:
            return
        self._word_section(document, labels.experience)
        for experience in content.experience:
            self._word_company(document, experience, labels)

    def _word_other_sections(
        self, *, document: WordDocument, content: ResumeContent, labels: ResumeExportLabels
    ) -> None:
        if content.certifications:
            self._word_section(document, labels.certifications)
            for certification in content.certifications:
                self._word_heading(document, f"{certification.name} | {certification.issuer}")
                dates = self._certification_dates(item=certification, language=labels.language)
                if dates:
                    document.add_paragraph(dates)
                if certification.credential_url:
                    document.add_paragraph(certification.credential_url)
        for extra in content.additional_sections:
            self._word_section(document, extra.title)
            for entry in extra.items:
                self._word_heading(document, entry.title)
                if entry.description:
                    document.add_paragraph(entry.description)
                if entry.url:
                    document.add_paragraph(entry.url)

    def _word_section(self, document: WordDocument, title: str) -> None:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_before = Pt(6)
        paragraph.paragraph_format.space_after = Pt(3)
        paragraph.paragraph_format.keep_with_next = True
        run = paragraph.add_run(title.upper())
        run.bold = True
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(37, 84, 135)
        borders = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "6")
        bottom.set(qn("w:color"), "255487")
        borders.append(bottom)
        paragraph._p.get_or_add_pPr().append(borders)  # noqa: SLF001

    def _word_heading(self, document: WordDocument, title: str) -> None:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.keep_with_next = True
        run = paragraph.add_run(title)
        run.bold = True
        run.font.size = Pt(9)

    def _word_band(self, document: WordDocument, text: str, *, blue: bool) -> None:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(1)
        run = paragraph.add_run(text)
        run.bold = blue
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), "F1F7FC" if blue else "F6F9FC")
        paragraph._p.get_or_add_pPr().append(shading)  # noqa: SLF001

    def _word_company(
        self, document: WordDocument, item: ResumeExperienceItem, labels: ResumeExportLabels
    ) -> None:
        period = self._format_date_range(
            start_date=item.start_date,
            end_date=item.end_date,
            current_status=item.current_status,
            labels=labels,
        )
        self._word_band(document, f"{item.company} | {period}", blue=True)
        document.add_paragraph(" | ".join(part for part in (item.position, item.location) if part))
        if item.summary:
            document.add_paragraph(item.summary)
        for highlight in item.highlights:
            document.add_paragraph(highlight, style="List Bullet")
        if item.technologies:
            self._word_band(
                document, f"{labels.technologies}: {', '.join(item.technologies)}", blue=False
            )
        for project in item.projects:
            project_title = f"{labels.project}: {project.name}"
            if project.role:
                project_title += f" | {project.role}"
            self._word_band(document, project_title, blue=True)
            for highlight in project.highlights:
                document.add_paragraph(highlight, style="List Bullet")
            if project.description:
                self._word_band(document, project.description, blue=False)
            if project.technologies:
                self._word_band(
                    document,
                    f"{labels.technologies}: {', '.join(project.technologies)}",
                    blue=False,
                )
            if project.url:
                document.add_paragraph(project.url)

    def _certification_dates(self, *, item: ResumeCertificationItem, language: LanguageEnum) -> str:
        parts = []
        if item.issued_on:
            label = "Выдан" if language == LanguageEnum.RU else "Issued"
            parts.append(f"{label}: {self._format_date(value=item.issued_on, language=language)}")
        if item.expires_on:
            label = "Истекает" if language == LanguageEnum.RU else "Expires"
            parts.append(f"{label}: {self._format_date(value=item.expires_on, language=language)}")
        return " | ".join(parts)
